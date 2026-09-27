import docx
import sys
import os

sys.stdout.reconfigure(encoding="utf-8")
docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final.docx")
doc = docx.Document(docx_path)

for i, p in enumerate(doc.paragraphs):
    if "图 3" in p.text:
        print(f"Figure 3 caption is at P{i}: {p.text}")
        if i > 0:
            prev_p = doc.paragraphs[i-1]
            has_drawing = "drawing" in prev_p._element.xml
            print(f"  P{i-1} text: {repr(prev_p.text)}, has drawing: {has_drawing}")
        if i + 1 < len(doc.paragraphs):
            next_p = doc.paragraphs[i+1]
            print(f"  P{i+1} text: {repr(next_p.text)}")
