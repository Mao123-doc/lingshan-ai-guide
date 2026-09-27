from __future__ import annotations

import sys
from pathlib import Path

VENDOR = Path(__file__).with_name("vendor")
sys.path.insert(0, str(VENDOR))

from pptx import Presentation  # type: ignore  # noqa: E402


def slide_text(slide) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text:
            parts.append(shape.text)
    return "\n".join(parts)


def main() -> int:
    if len(sys.argv) != 2:
        raise SystemExit("usage: verify_pptx.py <deck.pptx>")
    deck_path = Path(sys.argv[1])
    prs = Presentation(deck_path)
    failures: list[str] = []

    if len(prs.slides) != 14:
        failures.append(f"expected 14 slides, found {len(prs.slides)}")
    expected_w = 10 * 914400
    expected_h = 5.625 * 914400
    if abs(prs.slide_width - expected_w) > 10 or abs(prs.slide_height - expected_h) > 10:
        failures.append(f"wrong canvas: {prs.slide_width} x {prs.slide_height}")

    critical = {
        4: ["23 景点", "31 道路", "14 项守恒"],
        5: ["Promise.all", "RRF 融合", "k = 60", "Cross-Encoder"],
        7: ["98.0%", "96.0%", "48 / 50"],
        8: ["FORBIDDEN_LLM_FIELDS", "≥ 0.75", "40 / 40"],
        9: ["Dijkstra", "Beam Search", "15–20 分钟"],
        10: ["96.0%", "40 / 40", "60 / 60", "物理硬违规 0 次"],
        11: ["诚实拒绝", "只剩 20 分钟"],
        12: ["200 QPS级", "证据追溯", "实景导航"],
        13: ["< 1.5 万元", "3–6 个月", "1周", "2周", "4周"],
    }
    for slide_no, terms in critical.items():
        text = slide_text(prs.slides[slide_no - 1])
        for term in terms:
            if term not in text:
                failures.append(f"slide {slide_no}: missing {term!r}")

    sw, sh = prs.slide_width, prs.slide_height
    for slide_no, slide in enumerate(prs.slides, start=1):
        if slide_no > 1 and f"{slide_no:02d}" not in slide_text(slide):
            failures.append(f"slide {slide_no}: missing page badge")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > sw + 10 or shape.top + shape.height > sh + 10:
                failures.append(f"slide {slide_no}: shape {shape.shape_id} outside canvas")

    if failures:
        print("FAIL")
        for failure in failures:
            print(f"- {failure}")
        return 1
    print(f"PASS: {deck_path}")
    print("slides=14 canvas=10x5.625in page_badges=13 critical_claims=ok bounds=ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
