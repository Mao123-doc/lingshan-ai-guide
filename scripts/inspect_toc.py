import docx
import sys
import os

sys.stdout.reconfigure(encoding="utf-8")
docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final.docx")
doc = docx.Document(docx_path)

print("=== PARAGRAPHS 0 to 55 ===")
for i in range(min(55, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    print(f"P{i:02d} [{p.style.name}]: {repr(p.text)}")
