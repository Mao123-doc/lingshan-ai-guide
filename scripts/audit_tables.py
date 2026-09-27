import docx
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final.docx")
doc = docx.Document(docx_path)

print(f"=== TABLES DETAILED FORMATTING AUDIT ({len(doc.tables)} tables) ===")

for idx, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns)
    
    # Check alignment of table
    alignment = tbl.alignment
    
    # Check header repeat (tblHeader)
    header_tr = tbl.rows[0]._tr.get_or_add_trPr()
    tbl_headers = header_tr.xpath(".//w:tblHeader")
    cant_splits = header_tr.xpath(".//w:cantSplit")
    
    # Check cell text length / preview
    preview = []
    if rows > 0:
        for c in tbl.rows[0].cells[:min(4, cols)]:
            preview.append(c.text.strip().replace("\n", " "))
            
    has_header_repeat = len(tbl_headers) > 0
    has_cant_split = len(cant_splits) > 0
    
    print(f"Table {idx+1:02d}: {rows} rows x {cols} cols | align={alignment} | tblHeader={has_header_repeat} | cantSplit={has_cant_split} | cols: {preview}")
