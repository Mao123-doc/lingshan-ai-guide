import docx
import os

src = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final.docx")
dst = os.path.join(os.getcwd(), "tmp", "test_delete_res.docx")

doc = docx.Document(src)
print("Initial paragraphs:", len(doc.paragraphs))

# Find paragraph with "路线步骤"
for p in list(doc.paragraphs):
    if p.text.strip() == "路线步骤":
        print("Found '路线步骤', deleting...")
        el = p._element
        el.getparent().remove(el)

doc.save(dst)

doc2 = docx.Document(dst)
print("Saved paragraphs:", len(doc2.paragraphs))
print("Successfully saved and reloaded!")
