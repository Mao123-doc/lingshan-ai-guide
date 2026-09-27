import docx
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8")

docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final.docx")
backup_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final_backup_prepatch.docx")
img12_path = os.path.join(os.getcwd(), "tmp", "image12.jpeg")

print(f"Target file: {docx_path}")
print(f"Backup destination: {backup_path}")

# 1. Create safety backup
shutil.copy2(docx_path, backup_path)
print("Backup created successfully.")

# 2. Open document
doc = docx.Document(docx_path)
print(f"Loaded document. Total paragraphs: {len(doc.paragraphs)}")

# Helper to safely delete a paragraph
def delete_paragraph(p):
    el = p._element
    el.getparent().remove(el)

# Step A: Insert missing picture for 图 8 (诚实拒绝与降级微游)
target_cap = None
for p in doc.paragraphs:
    if "极限约束下的诚实拒绝与就近降级" in p.text:
        target_cap = p
        break

if target_cap is not None:
    print("Found caption for 图 8, inserting picture before it...")
    pic_p = target_cap.insert_paragraph_before()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.keep_with_next = True
    pic_p.paragraph_format.space_before = Pt(8)
    pic_p.paragraph_format.space_after = Pt(4)
    run = pic_p.add_run()
    run.add_picture(img12_path, width=Cm(12.0))
    print("Successfully inserted image12.jpeg!")
else:
    print("ERROR: Caption for 图 8 not found!")

# Step B: Update figure captions
for p in doc.paragraphs:
    txt = p.text.strip()
    
    # 图 4- 1 -> 图 5
    if txt.startswith("图 4- 1") or txt.startswith("图 4-1"):
        print(f"Updating 图 4- 1 -> 图 5 (runs={len(p.runs)})")
        p.runs[2].text = "5  "
        p.runs[3].text = ""
        p.runs[4].text = ""
        # Append '支' if needed
        if not p.text.endswith("支"):
            p.runs[-1].text += "支"
            
    # 图 7- 1 -> 图 8
    elif txt.startswith("图 7- 1") or txt.startswith("图 7-1"):
        print(f"Updating 图 7- 1 -> 图 8 (runs={len(p.runs)})")
        p.runs[2].text = "8  "
        p.runs[3].text = ""
        p.runs[4].text = ""

    # 图 8 -> 图 9
    elif txt.startswith("图 8  极限约束"):
        print(f"Updating 图 8 -> 图 9 (runs={len(p.runs)})")
        p.runs[2].text = "9  "
        if not p.text.endswith("览"):
            p.runs[-1].text += "览"

    # 图 9 -> 图 10
    elif txt.startswith("图 9  问答页"):
        print(f"Updating 图 9 -> 图 10 (runs={len(p.runs)})")
        p.runs[2].text = "10"

    # 图 10 -> 图 11
    elif txt.startswith("图 10  语音交互"):
        print(f"Updating 图 10 -> 图 11 (runs={len(p.runs)})")
        p.runs[2].text = "11  "
        if not p.text.endswith("应"):
            p.runs[-1].text += "应"

    # 图 11 -> 图 12
    elif txt.startswith("图 11  知识库"):
        print(f"Updating 图 11 -> 图 12 (runs={len(p.runs)})")
        p.runs[2].text = "12  "

    # 图 12 -> 图 13
    elif txt.startswith("图 12  运营数据"):
        print(f"Updating 图 12 -> 图 13 (runs={len(p.runs)})")
        p.runs[2].text = "13  "

    # 图 13 -> 图 14
    elif txt.startswith("图 13  七组检索"):
        print(f"Updating 图 13 -> 图 14 (runs={len(p.runs)})")
        p.runs[2].text = "14  "

    # 图 14 -> 图 15
    elif txt.startswith("图 14  边缘案例一"):
        print(f"Updating 图 14 -> 图 15 (runs={len(p.runs)})")
        p.runs[2].text = "15  "
        if not p.text.endswith("(PASS)"):
            p.runs[-1].text += "(PASS)"

    # 图 15 -> 图 16
    elif txt.startswith("图 15  边缘案例二"):
        print(f"Updating 图 15 -> 图 16 (runs={len(p.runs)})")
        p.runs[2].text = "16  "
        if not p.text.endswith("(PASS)"):
            p.runs[-1].text += " 1.00(PASS)"

    # 图 16 -> 图 17
    elif txt.startswith("图 16  原型路线"):
        print(f"Updating 图 16 -> 图 17 (runs={len(p.runs)})")
        p.runs[2].text = "17  "

    # 图 17 -> 图 18
    elif txt.startswith("图 17  原型需求"):
        print(f"Updating 图 17 -> 图 18 (runs={len(p.runs)})")
        p.runs[2].text = "18  "

    # 图 18 -> 图 19
    elif txt.startswith("图 18  信息不足"):
        print(f"Updating 图 18 -> 图 19 (runs={len(p.runs)})")
        p.runs[2].text = "19  "

# Step C: Update in-text citations
for p in doc.paragraphs:
    # 1. 图 8 展示受限预算下的拒绝案例
    if "图 8 展示受限预算下的拒绝案例" in p.text or ("图 8" in p.text and "受限预算下的拒绝案例" in p.text):
        print("Updating in-text citation: 图 8 展示 -> 图 9 展示")
        for r in p.runs:
            if r.text == "8 ":
                r.text = "9 "
                break

    # 2. 核心 AI 能力被封装为可交互的 Web 产品：...（图 9）...（图 10）
    if "核心 AI 能力被封装为可交互的 Web 产品" in p.text:
        print("Updating in-text citations in 3.3.8: （图 9）->（图 10）, （图 10）->（图 11）")
        for idx, r in enumerate(p.runs):
            if r.text == "9" and idx > 0 and "（图" in p.runs[idx-1].text:
                r.text = "10"
            elif r.text == "10" and idx > 0 and "（图" in p.runs[idx-1].text:
                r.text = "11"

    # 3. 前端交互：React 实现问答页...（图 11）
    if "知识库管理页支持文件索引状态查看、上传重建与 RAG 测试" in p.text:
        print("Updating in-text citation in 4.2: （图 11）->（图 12）")
        for idx, r in enumerate(p.runs):
            if r.text == "11" and idx > 0 and "（图" in p.runs[idx-1].text:
                r.text = "12"

    # 4. 特别说明：图 12 看板
    if "特别说明：图 12 看板" in p.text or ("特别说明：图" in p.text and "12 " in p.text):
        print("Updating in-text citation: 图 12 看板 -> 图 13 看板")
        for r in p.runs:
            if r.text == "12 ":
                r.text = "13 "
                break

    # 5. 在同一 50 题集上比较七种检索配置，结果见表 17 和图 13 。
    if "在同一 50 题集上比较七种检索配置" in p.text:
        print("Updating in-text citation: 和图 13 -> 和图 14")
        for idx, r in enumerate(p.runs):
            if r.text == "13" and idx > 0 and "图" in p.runs[idx-1].text:
                r.text = "14"
                break

    # 6. 见图 14 、图 15；
    if "见图 14 、图 15" in p.text or ("见图" in p.text and "14" in p.text and "15" in p.text):
        print("Updating in-text citations: 见图 14 、图 15 -> 见图 15 、图 16")
        for idx, r in enumerate(p.runs):
            if r.text == "14":
                r.text = "15"
            elif r.text == "15":
                r.text = "16"

    # 7. 而非直接猜测规划，如图 18。
    if "而非直接猜测规划，如图 18" in p.text or ("而非直接猜测规划" in p.text and "18" in p.text):
        print("Updating in-text citation: 如图 18 -> 如图 19")
        for idx, r in enumerate(p.runs):
            if r.text == "18":
                r.text = "19"
                break

# Step D: Delete broken orphan paragraphs
to_delete = []
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt in ["支", "路线步骤", "览", "览览", "应", "(PASS)", "1.00(PASS)"]:
        print(f"Marking orphan paragraph for deletion: '{txt}'")
        to_delete.append(p)

for p in to_delete:
    delete_paragraph(p)
print(f"Deleted {len(to_delete)} orphan paragraphs.")

# 3. Save document
doc.save(docx_path)
print("Saved patched document to final.docx successfully!")
