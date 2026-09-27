import docx
import sys
import os

sys.stdout.reconfigure(encoding="utf-8")
src = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final.docx")
doc = docx.Document(src)

for idx in range(min(179, len(doc.paragraphs))):
    p = doc.paragraphs[idx]
    txt = p.text.strip()
    if "图" in txt or (0 < len(txt) <= 5):
        print(f"[{idx:03d}] (len={len(txt)}): '{txt}'")
