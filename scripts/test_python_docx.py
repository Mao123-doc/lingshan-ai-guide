import docx
import os
import zipfile

src = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "upload", "final.docx")
dst = os.path.join(os.getcwd(), "tmp", "test_docx_save.docx")

doc = docx.Document(src)
print("Paragraphs in docx.Document:", len(doc.paragraphs))
print("Inline shapes in docx.Document:", len(doc.inline_shapes))
print("Tables in docx.Document:", len(doc.tables))

# Save without modifying
doc.save(dst)

# Check media files in saved docx
with zipfile.ZipFile(dst, "r") as z:
    media = [f for f in z.namelist() if f.startswith("word/media/")]
    print("Media files in saved docx:", len(media))
