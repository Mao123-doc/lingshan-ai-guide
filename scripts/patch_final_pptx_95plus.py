from pathlib import Path
import argparse

from pptx import Presentation
from pptx.util import Pt


DEFAULT_PPTX_PATH = Path(
    r"release\competition_final\01_汇报PPT与导出版\灵小禅_竞赛汇报PPT.pptx"
)


def replace_text(prs: Presentation, old: str, new: str) -> int:
    changed = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if old in shape.text:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old, new)
                            changed += 1
                if old in shape.text:
                    shape.text_frame.text = shape.text.replace(old, new)
                    changed += 1
    return changed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pptx", type=Path, default=DEFAULT_PPTX_PATH)
    args = parser.parse_args()
    pptx_path = args.pptx
    prs = Presentation(pptx_path)

    # P5: Router is explanatory metadata, not a selector that disables branches.
    changed = 0
    p5 = prs.slides[4]
    for shape in p5.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.text == "问题分析 / Router":
            shape.text_frame.text = "问题分析 / Query Analysis"
            changed += 1
        elif shape.text == "判断问题性质，决定检索通道":
            shape.text_frame.text = "识别问题类型，仅用于 Trace 标注与结果解释"
            changed += 1
        elif shape.text == "不同问题 → 不同通道（真实映射）":
            shape.text_frame.text = "不同问题 → 不同解释重点（三路并行召回）"
            changed += 1
        elif shape.text == "多路混合 RAG：让不同性质的问题走不同的信息通道":
            shape.text_frame.text = "多路混合 RAG：让不同性质的问题获得互补证据"
            changed += 1

    # Keep the final deck's claims inside the evidence boundary.
    changed += replace_text(prs, "从源头根除事实性幻觉", "降低事实性错误并保留证据追溯")
    changed += replace_text(prs, "知识零幻觉", "事实合约可追溯")
    changed += replace_text(prs, "硬约束 100% 精准履约，0 次物理违规！", "当前结构化约束测试中 0 次物理硬违规")
    changed += replace_text(prs, "总耗时精准 300 分钟 · 0 次物理硬违规！", "总耗时 300 分钟 · 当前结构化约束测试中 0 次物理硬违规")
    changed += replace_text(prs, "总耗时精准 300 分钟·0次物理硬违规!", "总耗时 300 分钟·当前结构化约束测试中 0 次物理硬违规")
    changed += replace_text(prs, "14 项形式化守恒仲裁全绿放行", "14 项形式化守恒仲裁通过检查")
    changed += replace_text(prs, "强劲商业闭环", "商业闭环（情景模型）")
    changed += replace_text(prs, "3~6 个月收回 ROI", "情景模型预计 3~6 个月收回投入")

    # Keep the editable deck aligned with the audited route asset ledger.
    changed += replace_text(
        prs,
        "▪ 31 条实测双向道路",
        "▪ 30 条当前登记道路边（31 条目标口径待补齐核验）",
    )
    changed += replace_text(
        prs,
        "【23 核心景点 + 31 实测道路数字孪生路网】",
        "【23 个路线节点 + 30 条当前登记道路边（31 条目标口径待补齐核验）】",
    )
    changed += replace_text(prs, "(23 POI + 31 道路)", "(23 节点 + 30 条登记边；31 条目标待核验)")
    changed += replace_text(prs, "23 个 POI · 31 条实测双向道路", "23 个路线节点 · 30 条当前登记道路边（31 条目标待核验）")
    changed += replace_text(prs, "23 个景点 POI + 31 条实测双向道路", "23 个路线节点 + 30 条当前登记道路边（31 条目标待核验）")

    # P11: widen the time-label text boxes and shorten long labels to prevent
    # overlap with the gold pills in the exported deck.
    p11 = prs.slides[10]
    p11_labels = {
        "12:30 突发超时": "12:30 动态纠偏",
        "13:40 抵达梵宫": "13:40 梵宫检票",
        "15:10 换乘无障碍电瓶车（约 0m 步行）": "15:10 无障碍电瓶车接驳",
    }
    for shape in p11.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.text in p11_labels:
            shape.text_frame.text = p11_labels[shape.text]
            changed += 1
        if shape.text in {"11:00 正门启程", "12:30 动态纠偏", "13:40 梵宫检票", "15:10 无障碍电瓶车接驳"}:
            shape.width = int(2.35 * 914400)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(15)
        if shape.text.startswith(("【极限防御实测 · 诚实拒绝与优雅降级】", "【诚实拒绝与优雅降级】")):
            shape.text_frame.text = (
                "【诚实拒绝与优雅降级】\n"
                "▪ 20 分钟无法同时完成正门往返与 14:00 演出，系统拒绝编造路线；\n"
                "▪ 自动剥离演艺偏好，生成正门广场与佛足坛微游览方案（18min），保留可执行替代。"
            )
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
            changed += 1
        elif shape.text.startswith("【主动澄清与追问机制】"):
            shape.text_frame.text = "主动澄清机制已在 P8 展示：缺失时间或位置时先追问，不猜测、不编造。"
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
            changed += 1

    # P2/P11/P13: make the evidence boundary and page hierarchy explicit.
    changed += replace_text(prs, "【探索性预调研 n=50】：", "【探索性预调研 n=50，不外推至全部游客】：")
    changed += replace_text(prs, "典型场景实测：限时 5 小时带老人看演出游大佛 2.5D 物理动线全还原", "典型场景实测：限时 5 小时带老人看演出游大佛 2.5D 物理动线复盘")
    changed += replace_text(prs, "推广应用价值：三阶轻量化落地体系与 3 年商业 ROI 测算模型", "推广应用价值：三阶轻量化落地与情景 ROI 模型")
    for shape in prs.slides[12].shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.text.startswith("3 年商业财务测算与 ROI 模型"):
            shape.text_frame.text = (
                "商业财务测算（情景假设）\n"
                "💵 年运维：1.39 万元；首次实施：5.0 万元（独立核算）\n"
                "📈 闭环：SaaS 服务费 + 个性化交付 + 二消/观光车分成\n"
                "🎯 ROI：3~6 个月回收投入（中性情景 4.8 个月；非已实现经营结果）"
            )
            changed += 1

    if len(prs.slides) != 14:
        raise SystemExit(f"Expected 14 slides, found {len(prs.slides)}")

    prs.save(pptx_path)
    print(f"Patched {pptx_path} ({changed} text changes); slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
