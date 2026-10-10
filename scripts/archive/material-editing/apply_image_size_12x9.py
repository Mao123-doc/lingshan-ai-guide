# -*- coding: utf-8 -*-
"""
scripts/archive/material-editing/apply_image_size_12x9.py
Adjusts all 18 figures to Height 9cm, Width 12cm (4320000 x 3240000 EMUs),
centers them, pairs them with captions via keep_with_next,
removes duplicate paragraph P139 before Figure 5 to prevent caption detachment,
and ensures zero layout disruptions across the document.
Leaves formula drawings (equations 3-1, 3-2, 3-7) completely untouched!
"""

import docx, win32com.client, os
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

DOC_PATH = r"E:\Projects\lingshan-ai-guide-Mao123-doc\竞赛汇报与文档材料包\03_申报文档与技术报告\Word工作稿\final.docx"

# 12cm = 12 * 360000 = 4320000 EMUs
# 9cm  =  9 * 360000 = 3240000 EMUs
TARGET_CX = 4320000
TARGET_CY = 3240000

def run():
    doc = docx.Document(DOC_PATH)
    print(f"Loaded {DOC_PATH}, paragraphs: {len(doc.paragraphs)}")

    # 1. Remove duplicate paragraph P139 before Figure 5
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t == "问答页提供可展开的证据抽屉，显示来源文件、知识片段与相关度等信息。":
            print(f"Removing duplicate drawer paragraph P{i}...")
            p._element.getparent().remove(p._element)
            break

    # 2. Locate all 18 figure image paragraphs and captions
    # Note: figures are paragraphs containing w:drawing that are immediately followed by a caption starting with '图 '
    adjusted_count = 0
    for i, p in enumerate(doc.paragraphs):
        if 'w:drawing' in p._element.xml:
            # check next paragraph for caption
            next_p = doc.paragraphs[i + 1] if i + 1 < len(doc.paragraphs) else None
            if next_p and next_p.text.strip().startswith('图 '):
                # This is a confirmed Figure (not a formula)
                cap_text = next_p.text.strip()[:40]
                print(f"Adjusting Figure at P{i} -> {cap_text} to 12cm x 9cm...")

                # Update wp:extent
                for node in p._element.xpath('.//wp:extent'):
                    node.attrib['cx'] = str(TARGET_CX)
                    node.attrib['cy'] = str(TARGET_CY)

                # Update a:xfrm/a:ext
                for node in p._element.xpath('.//a:xfrm/a:ext'):
                    node.attrib['cx'] = str(TARGET_CX)
                    node.attrib['cy'] = str(TARGET_CY)

                # Format image paragraph
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = 0
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = 38100 # 3pt
                p.paragraph_format.space_after = 25400  # 2pt

                # Format caption paragraph
                next_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                next_p.paragraph_format.first_line_indent = 0
                next_p.paragraph_format.space_before = 25400 # 2pt
                next_p.paragraph_format.space_after = 38100  # 3pt
                adjusted_count += 1

    print(f"Total figures adjusted to 12cm x 9cm: {adjusted_count} (expected 18)")

    doc.save(DOC_PATH)
    print("Document saved successfully!")

    # 3. Verify with Word COM
    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc_word = word.Documents.Open(os.path.abspath(DOC_PATH))
        pages = doc_word.ComputeStatistics(2)
        words = doc_word.ComputeStatistics(0)
        print(f"\n=== Word COM Verification ===")
        print(f"Total Pages: {pages}, Total Words: {words}")

        # Check all 18 figures for page alignment
        detached = 0
        for fig_num in range(1, 19):
            search_text = f"图 {fig_num} "
            rng = doc_word.Content
            if rng.Find.Execute(search_text):
                cap_page = rng.Information(3) # wdActiveEndPageNumber
                # check image page by inspecting range immediately preceding the caption
                img_rng = rng.Duplicate
                img_rng.Collapse(1) # start
                img_rng.MoveStart(4, -1) # move back 1 paragraph
                img_page = img_rng.Information(3)
                is_ok = (cap_page == img_page)
                if not is_ok:
                    detached += 1
                    print(f"  WARNING: Figure {fig_num} MISMATCH! Img Page={img_page}, Cap Page={cap_page}")
                else:
                    print(f"  Figure {fig_num:>2}: Page {cap_page} (Aligned OK)")

        print(f"\nCaption Detachment Check: {detached} mismatches (Goal: 0)")
        doc_word.Close(False)
    finally:
        word.Quit()

if __name__ == '__main__':
    run()
