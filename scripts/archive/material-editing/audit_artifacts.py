import docx
import os
import sys
import re

sys.stdout.reconfigure(encoding="utf-8")

docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "03_申报文档与技术报告", "Word工作稿", "final.docx")
doc = docx.Document(docx_path)

print("=== 1. FORMULA & OCR ARTIFACTS ===")
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if "__" in txt or re.search(r"l2\s*个", txt) or re.search(r"\(\s*3\s*__", txt) or "ToP" in txt or "ExPand" in txt or re.search(r"ei\s*__", txt):
        print(f"P{idx:03d}: {txt}")

print("\n=== 2. TRUNCATED / BROKEN SENTENCES ===")
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    # Check if ends with dangling punctuation
    if txt.endswith("、") or txt.endswith("，") or txt.endswith("；") or txt.endswith("：") or txt.endswith("（") or txt.endswith("＋") or txt.endswith("+"):
        if len(txt) > 10 and not p.style.name.startswith("Heading"):
            print(f"P{idx:03d} (ends with dangling punct): {txt}")
    elif 0 < len(txt) < 8 and not txt.startswith("第") and not txt.startswith("表") and not txt.startswith("图"):
        if p.style.name != "Heading 1" and p.style.name != "Heading 2":
            print(f"P{idx:03d} (very short): '{txt}'")
