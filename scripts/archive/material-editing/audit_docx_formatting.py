import docx
from docx.oxml import OxmlElement
import os
import sys
import re

sys.stdout.reconfigure(encoding="utf-8")

docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "03_申报文档与技术报告", "Word工作稿", "final.docx")
doc = docx.Document(docx_path)

print(f"=== DOCX FORMATTING AUDIT ===")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# 1. Formula & artifact scan
print("\n--- 1. Searching for OCR/Formula Artifacts (__ or l2 or double underscore) ---")
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if "__" in txt or re.search(r"l2\s*个", txt) or re.search(r"\(\s*3\s*__", txt) or "ToP" in txt or "ExPand" in txt:
        print(f"P{idx:03d}: {txt}")

# 2. Broken sentences (ending with comma, semicolon, dash, or dangling phrase)
print("\n--- 2. Searching for Truncated/Broken Sentences ---")
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    # Check if ends with dangling punctuation
    if txt.endswith("、") or txt.endswith("，") or txt.endswith("；") or txt.endswith("：") or txt.endswith("（") or txt.endswith("＋") or txt.endswith("+"):
        # filter out table or heading
        if len(txt) > 10 and not p.style.name.startswith("Heading"):
            print(f"P{idx:03d} (ends with dangling punct): {txt}")
    # Check if very short paragraph that is not heading or table
    elif 0 < len(txt) < 8 and not txt.startswith("第") and not txt.startswith("表") and not txt.startswith("图"):
        if p.style.name != "Heading 1" and p.style.name != "Heading 2":
            print(f"P{idx:03d} (very short): '{txt}'")

# 3. Table captions and numbering scan
print("\n--- 3. Table Captions and Numbering Scan ---")
table_caps = []
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if re.match(r"^表\s*\d+", txt):
        table_caps.append((idx, txt))
        print(f"P{idx:03d}: {txt}")

# 4. Heading styles and numbering check
print("\n--- 4. Heading Hierarchy & Numbering Check ---")
headings = []
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name.startswith("Heading") or re.match(r"^\d+(\.\d+)*\s+", txt) or re.match(r"^第[一二三四五六七八九十]+章", txt):
        headings.append((idx, p.style.name, txt))
        print(f"P{idx:03d} [{p.style.name}]: {txt}")

# 5. Consecutive blank lines check
print("\n--- 5. Consecutive Empty Paragraphs Scan ---")
empty_streak = 0
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    drawings = len(p._element.xpath(".//w:drawing"))
    if not txt and drawings == 0:
        empty_streak += 1
        if empty_streak >= 3:
            print(f"P{idx:03d}: Streak of {empty_streak} consecutive empty paragraphs")
    else:
        empty_streak = 0
