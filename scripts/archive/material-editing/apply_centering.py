import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import shutil
import sys
import re

sys.stdout.reconfigure(encoding="utf-8")

docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "03_申报文档与技术报告", "Word工作稿", "final.docx")
backup_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "03_申报文档与技术报告", "Word工作稿", "final_backup_pre_centering.docx")

print(f"Target file: {docx_path}")
print(f"Backup file: {backup_path}")

shutil.copy2(docx_path, backup_path)
print("Backup created successfully.")

doc = docx.Document(docx_path)
print(f"Loaded document. Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")

def delete_paragraph(p):
    el = p._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)

def has_drawing(p):
    return len(p._element.xpath(".//*[local-name()='drawing']")) > 0

def set_perfect_center(p, space_before=Pt(6), space_after=Pt(6), keep_with_next=False):
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.right_indent = Pt(0)
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    p.paragraph_format.keep_with_next = keep_with_next

# ==========================================================
# 1. Centering Table Names (表格的名) & Remove Spurious Empty Paragraphs
# ==========================================================
print("\n--- 1. Centering Table Captions ---")
tbl_captions_count = 0
empty_before_table_removed = 0

for t_idx, tbl in enumerate(doc.tables):
    prev_el = tbl._element.getprevious()
    # Find paragraph immediately preceding tbl
    p_between = None
    caption_p = None
    
    # Check if directly preceding is an empty paragraph
    if prev_el is not None and prev_el.tag.endswith('p'):
        matching_ps = [p for p in doc.paragraphs if p._element == prev_el]
        if matching_ps:
            p_cand = matching_ps[0]
            txt = p_cand.text.strip()
            if not txt and not has_drawing(p_cand):
                p_between = p_cand
                # Look one element higher for caption
                above_el = prev_el.getprevious()
                if above_el is not None and above_el.tag.endswith('p'):
                    matching_above = [p for p in doc.paragraphs if p._element == above_el]
                    if matching_above and re.match(r"^表\s+\d+\s+", matching_above[0].text.strip()):
                        caption_p = matching_above[0]
            elif re.match(r"^表\s+\d+\s+", txt):
                caption_p = p_cand
                
    if caption_p is not None:
        tbl_captions_count += 1
        set_perfect_center(caption_p, space_before=Pt(8), space_after=Pt(4), keep_with_next=True)
        print(f"Table {t_idx+1:02d} Caption: '{caption_p.text[:35]}...' -> Centered & keep_with_next=True")
        
        # Remove empty paragraph between caption and table if present
        if p_between is not None:
            delete_paragraph(p_between)
            empty_before_table_removed += 1

print(f"Centered {tbl_captions_count}/20 table captions. Removed {empty_before_table_removed} redundant empty lines between captions and tables.")

# ==========================================================
# 2. Centering Figure Names (图片的名) & Figure Images
# ==========================================================
print("\n--- 2. Centering Figure Images and Captions ---")
fig_count = 0

for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    # Check for figure caption
    if re.match(r"^图\s+\d+\s+", txt) and len(txt) < 80 and not txt.endswith("。") and not txt.endswith("；"):
        fig_count += 1
        set_perfect_center(p, space_before=Pt(2), space_after=Pt(8), keep_with_next=False)
        print(f"Figure {fig_count:02d} Caption [P{idx:03d}]: '{txt[:35]}...' -> Centered")
        
        # Preceding paragraph must be the image
        if idx > 0 and has_drawing(doc.paragraphs[idx-1]):
            img_p = doc.paragraphs[idx-1]
            set_perfect_center(img_p, space_before=Pt(8), space_after=Pt(2), keep_with_next=True)
            print(f"  Figure {fig_count:02d} Image [P{idx-1:03d}] -> Centered & keep_with_next=True")

print(f"Centered {fig_count}/18 figure captions and their images.")

# ==========================================================
# 3. Centering Formulas (公式)
# ==========================================================
print("\n--- 3. Centering Formulas ---")
formula_count = 0

for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    has_dw = has_drawing(p)
    
    is_formula = False
    # Check for formula tags or drawing formulas
    if any(tag in txt for tag in ['(3-3)', '(3-4)', '(3-5)', '(3-6)', '(3-8)', '(3-9)', '(3-10)', '(3 - 3)', '(3 - 4)', '(3 - 5)', '(3 - 6)', '(3 - 8)', '(3 - 9)', '(3 - 10)']):
        is_formula = True
    elif has_dw and (txt == 'sim' or idx in [146, 147, 148, 170, 171, 172]):
        # Check next paragraph to confirm it's an equation paragraph
        if idx + 1 < len(doc.paragraphs) and any(kw in doc.paragraphs[idx+1].text for kw in ['式（3-2）', '式（3-7）', '设检索通道集合']):
            is_formula = True
            
    if is_formula:
        formula_count += 1
        set_perfect_center(p, space_before=Pt(6), space_after=Pt(6), keep_with_next=False)
        print(f"Formula {formula_count:02d} [P{idx:03d}, dw={has_dw}]: '{txt[:40]}' -> Centered")

print(f"Centered {formula_count}/10 formulas.")

# ==========================================================
# 4. Fix Appendix F Typo & Heading 2 on Appendices
# ==========================================================
print("\n--- 4. Fix Appendix F Typo & Heading Styles ---")
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("录 F  参考文献") or txt.startswith("录 F 参考文献"):
        p.text = "附录 F  参考文献"
        p.style = "Heading 2"
        p.paragraph_format.alignment = None
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.right_indent = Pt(0)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = "黑体"
            r.font.size = Pt(14)
            r.font.bold = True
        print("Fixed '录 F  参考文献' -> '附录 F  参考文献' and applied Heading 2!")
    elif re.match(r"^附录\s+[A-G]\s+", txt) and len(txt) < 30:
        p.style = "Heading 2"
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.right_indent = Pt(0)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        print(f"Applied Heading 2 to: '{txt}'")

# Save document
doc.save(docx_path)
print(f"\nSaved centered document to {docx_path} successfully!")
