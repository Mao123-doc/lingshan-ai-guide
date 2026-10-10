import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('竞赛汇报与文档材料包/03_申报文档与技术报告/Word工作稿/final.docx')
print('Auditing', len(doc.paragraphs), 'paragraphs and', len(doc.tables), 'tables...')

# 1. Check broken punctuation flow
for i in range(1, len(doc.paragraphs)):
    prev = doc.paragraphs[i-1].text.strip()
    curr = doc.paragraphs[i].text.strip()
    if not curr or not prev:
        continue
    if curr[0] in [';', '；', ',', '，', ')', '）', '、', '.', '。', '?', '？', '!', '！', '”', '’']:
        print(f'Broken punctuation flow P{i-1} -> P{i}:')
        print(f'   Prev: {prev[-25:]}')
        print(f'   Curr: {curr[:25]}')

# 2. Check unfinished sentences
end_punct = ['.', '。', '!', '！', '?', '？', ':', '：', '”', '’', ')', '）', '—', '；', ';', '、']
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt or p.style.name.startswith('Heading') or txt.startswith('表 ') or txt.startswith('图 '):
        continue
    if txt[-1] not in end_punct and len(txt) > 20:
        next_txt = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ''
        if next_txt and not next_txt.startswith('表') and not next_txt.startswith('图'):
            print(f'Unfinished sentence P{i}: ...{txt[-30:]}')
