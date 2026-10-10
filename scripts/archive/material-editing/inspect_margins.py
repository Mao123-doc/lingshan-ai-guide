import docx
import sys
import os

sys.stdout.reconfigure(encoding="utf-8")
docx_path = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "03_申报文档与技术报告", "Word工作稿", "final.docx")
doc = docx.Document(docx_path)

print("=== SECTIONS & MARGINS ===")
for idx, sec in enumerate(doc.sections):
    print(f"Section {idx+1}: page_width={sec.page_width.cm:.1f}cm, page_height={sec.page_height.cm:.1f}cm")
    print(f"  margins: top={sec.top_margin.cm:.1f}cm, bottom={sec.bottom_margin.cm:.1f}cm, left={sec.left_margin.cm:.1f}cm, right={sec.right_margin.cm:.1f}cm")

print("\n=== PARAGRAPH FORMAT SAMPLES ===")
# Sample 10 body paragraphs
sample_indices = [53, 54, 87, 88, 134, 150, 158, 255, 302, 350]
for idx in sample_indices:
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        pf = p.paragraph_format
        first_line = pf.first_line_indent.pt if pf.first_line_indent else 0
        line_spacing = pf.line_spacing
        space_before = pf.space_before.pt if pf.space_before else 0
        space_after = pf.space_after.pt if pf.space_after else 0
        fonts = set()
        sizes = set()
        for r in p.runs:
            if r.font.name:
                fonts.add(r.font.name)
            if r.font.size:
                sizes.add(r.font.size.pt)
        print(f"P{idx:03d}: first_line={first_line}pt | line_spacing={line_spacing} | before={space_before}pt, after={space_after}pt | fonts={fonts} | sizes={sizes} | txt={p.text[:30]}...")
