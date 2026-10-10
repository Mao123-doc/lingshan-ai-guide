import docx
import sys
import os

sys.stdout.reconfigure(encoding="utf-8")
src = os.path.join(os.getcwd(), "竞赛汇报与文档材料包", "03_申报文档与技术报告", "Word工作稿", "final.docx")
doc = docx.Document(src)

print(f"Total doc.paragraphs: {len(doc.paragraphs)}")

for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if "图" in txt or (0 < len(txt) <= 5):
        print(f"[{idx:03d}] (len={len(txt)}): '{txt}'")
