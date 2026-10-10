# -*- coding: utf-8 -*-
"""
scripts/archive/material-editing/streamline_docx_phase2.py
Phase 2 precision trimming: compresses remaining over-long paragraphs
(P97, P222, P276, P282, P286, P301, P302, App D) by ~2,100 chars,
achieving a total ~18% reduction and landing the document cleanly at 28-30 pages.
"""

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

DOC_PATH = r"E:\Projects\lingshan-ai-guide-Mao123-doc\竞赛汇报与文档材料包\03_申报文档与技术报告\Word工作稿\final.docx"

def set_para_text(p, text, first_line_indent=304800, line_spacing=1.25, space_after=38100):
    p.text = ''
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:eastAsia'), '宋体')
    p.paragraph_format.first_line_indent = first_line_indent
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = space_after

def run():
    doc = docx.Document(DOC_PATH)
    initial_chars = sum(len(p.text.strip()) for p in doc.paragraphs)
    print(f"Loaded {DOC_PATH}, initial chars: {initial_chars}")

    # 1. P97 in Chapter 2: Quantitative Survey Breakdown
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('表 2 和图 2 汇总了 50 份有效受访样本在四个核心问题上的回答分布'):
            new_p97 = (
                "表 2 和图 2 汇总了 50 份有效受访样本的统计结果，四类现实痛点高度凸显："
                "（1）时空演艺错配（Q1）：82.0%（41/50）的受访者曾因耗时误判或提前停止检票而错过演艺；"
                "（2）物理拓扑盲区（Q2）：全量样本中 66.0%（33/50）曾遭遇台阶陡坡折返，而在长辈及推车等行动受限群体中该比例高达 93.3%（28/30）；"
                "（3）垂直事实幻觉（Q3）：74.0%（37/50）遭遇过通用 AI 编造票价或耗时失准；"
                "（4）决策型导览意向（Q4）：90.0%（45/50）表达了对能自动避障与动态纠偏决策导览的强烈使用期待。"
            )
            set_para_text(p, new_p97)
            print("Streamlined: Chapter 2 Survey Breakdown P97")
            break

    # 2. P222 in Chapter 4: Deliverables verification mechanism
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('交付成果实行四级分类管理与严谨核验机制：'):
            new_p222 = (
                "交付成果实行四级分类核验机制：一级为核心算法（13 维状态抽取、三路 RAG、RRF 融合 k=60、重排、24 宽束搜索、14 项形式化仲裁器），"
                "均通过自动化单元测试与基准集验证；二级为工程原型与双屏交互（问答抽屉、路线看板、调起高德导航、Live2D 数字人），录屏完整存档；"
                "三级为待实地核验项（道路坡度米数仪器测绘、出园路径闭环、演出停止检票对接及 SUS 评测）；"
                "四级为后续规划项（实时客流感知与接驳车调度联动）。"
            )
            set_para_text(p, new_p222)
            print("Streamlined: Chapter 4 Deliverables Verification P222")
            break

    # 3. P276 in Chapter 6: Hero Case Input & Slots
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('针对典型的“带长辈、看演出、低负荷”复合时空约束场景'):
            new_p276 = (
                "针对典型的“带长辈、看演出、低负荷”复合场景，系统在准生产测试环境下（Git SHA: 0522c386，录屏 recordings/recording.webm）开展端到端求解验证。"
                "测试设定游客口语输入：“现在上午 11 点，我在正门南门，带腿脚不便的妈妈，到下午四点前有 5 个小时，想看两点《吉祥颂》，大佛一定要去，尽量少走路。”"
                "状态抽取模块精准解析出 13 维状态：出发点 south_gate、时刻 11:00、时长 300 分钟（死线 16:00）、行动受限 limited（激活无障碍路网）、必去景点 LS-008、指定演出 14:00 场次《吉祥颂》。"
            )
            set_para_text(p, new_p276)
            print("Streamlined: Chapter 6 Hero Case Input P276")
            break

    # 4. P282 in Chapter 6: Hero Case 4 Physical Boundaries
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('为确保系统在走向实地部署时具备可靠的物理可行性'):
            new_p282 = (
                "为保证实地可用性，项目组客观界定当前路线的现实物理边界并制定闭环措施："
                "（1）返程出园路径闭环：末站 15:38 结束，曼飞龙塔至南门约 1200 米（长辈步行需 31 分钟），"
                "计入返程预计 16:09 出园（超出死线 9 分钟），后续工程需将返程出口设为终点硬约束并触发动态剪枝；"
                "（2）演出停检缓冲：需将剧场提前 15 分钟停检与安检排队硬扣减；"
                "（3）坡度实测：目前为定性标签，实地需通过仪器采集精确坡度（<2.5°）与台阶步数；"
                "（4）接驳车网络：景区电瓶车班次与换乘网络正进一步拓扑建模。"
            )
            set_para_text(p, new_p282)
            print("Streamlined: Chapter 6 Hero Boundaries P282")
            break

    # 5. P286 in Chapter 6: Field Survey Protocol
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('为在实地部署阶段彻底闭环上述现实约束'):
            new_p286 = (
                "为实地闭环上述约束，项目组已制定现场走查方案："
                "（1）工具：采用手持激光测距仪、高精度气压坡度计与 GPS 穿戴设备；"
                "（2）采样动线：选定长辈陪游、轮椅避障、临界赶场与微游览 4 条典型路线；"
                "（3）实测指标：逐段采集物理米数、通行耗时、阶梯步数、纵坡坡度与出园耗时；"
                "（4）算法闭环：路网补齐返程边并增设 returnToExit 强约束，前置扣减 15 分钟停检缓冲，"
                "返程超时自动剪枝非必选景点，确保全程收敛在死线之内。"
            )
            set_para_text(p, new_p286)
            print("Streamlined: Chapter 6 Field Survey P286")
            break

    # 6. P301 in Chapter 6: Hardware & API
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('（2）硬件规格与 API 预算假设：'):
            new_p301 = (
                "（2）硬件规格与 API 预算：采用 4 核 8G 本地边缘工控机部署向量库与图运筹算法，拓扑寻路时延 < 20ms，保障游客轨迹不出园。"
                "API 预算按试运行日均 200 次问答测算（单次约 2k tokens，年调用 7.3 万次，消耗 1.46 亿 tokens），"
                "按 DeepSeek 官方刊例单价（综合折合约 1.5 元/100万 tokens）计算理论年 API 费用约 219 元，"
                "预留冗余后列支约 400 元/年（0.04 万元/年），实际商用需动态校准。"
            )
            set_para_text(p, new_p301)
            print("Streamlined: Chapter 6 Hardware & API P301")
            break

    # 7. P302 in Chapter 6: ROI Scenario Models
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('（3）商业回报（ROI）情景推演模型：'):
            new_p302 = (
                "（3）商业回报（ROI）情景推演：针对未来落地构建三类模型："
                "一为 ToC 适老定制导览，年客流 200 万人次按 0.5% 转化率收取 5 元无障碍导览服务费，年增收 5.0 万元，首年总投入 6.39 万元的静态回收期约 1.28 年（15 个月）；"
                "二为 ToB 咨询分流，日均承接 1,000 次咨询并分流 15%~20% 高频问询，相当于替代 1 名季节性导游（年成本约 4~6 万元），12~18 个月对冲投入；"
                "三为旺季文创二消提成（客单提 5 元），理想情景下回本期可压缩至 3~6 个月。"
            )
            set_para_text(p, new_p302)
            print("Streamlined: Chapter 6 ROI Models P302")
            break

    # 8. Appendix D: User Survey Data
    appd_p1 = None
    appd_p2 = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('定量记录：材料包“05_现场演示与用户调研/pre_survey_data.json”含 50 条匿名记录'):
            appd_p1 = p
        elif t.startswith('定性材料：包含受访者 A～D 的跨景区初步半结构化访谈'):
            appd_p2 = p

    if appd_p1 and appd_p2:
        new_appd = (
            "用户调研原件归档于材料包“05_现场演示与用户调研/”：定量数据 pre_survey_data.json 包含 50 条匿名结构化记录（P01～P50，涵盖渠道、人群画像、Q1~Q4 选项与文字备注）；"
            "定性访谈详见 pre_survey_summary.md，完整收录受访者 A~D 跨景区初探记录，以及灵山胜境现场案例 T-07（轮椅登梯受阻）与 T-19（提前停检跑空）的访谈实录与约束提炼。"
        )
        set_para_text(appd_p1, new_appd)
        appd_p2.text = ''
        print("Streamlined: Appendix D merged into 1 paragraph")

    # Clean empty paragraphs
    paras_to_remove = []
    for i, p in enumerate(doc.paragraphs):
        if not p.text.strip() and 'w:drawing' not in p._element.xml and i > 40:
            paras_to_remove.append(p)

    for p in paras_to_remove:
        p._element.getparent().remove(p._element)

    doc.save(DOC_PATH)
    final_chars = sum(len(p.text.strip()) for p in doc.paragraphs)
    total_reduction = (26564 - final_chars) / 26564 * 100
    print(f"Phase 2 complete! Final chars: {final_chars}, Total reduction from original: {26564 - final_chars} chars ({total_reduction:.2f}%)")

if __name__ == '__main__':
    run()
