#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
audit_competition_materials.py
--------------------------------------------------
自动化交叉审计门禁脚本：
针对灵小禅全国总决赛材料执行严格的红线检查与事实台账核验。
支持扫描 .md, .tex, .pptx, .pdf 文件。

核验红线：
1. 严禁虚构数据：98.3%, 97.5%
2. 路网规模一致：当前资产为 23 个路线节点 + 30 条登记道路边，31 条为待补齐核验目标；禁止把 31 条写成已完成资产
3. 故事线页数：严禁 13 页（必须为 14 页标准版）
4. 英雄案例时间线：严禁 13:00 / 180分钟 / 3小时 / 13:45（必须为 11:00 出发、300分钟/5小时、13:40检票口带20分缓冲）
5. 98% 观察值界定：严禁将 98% 称为基线（基线必须为 96.0%，98% 必须加注消融/观察值）
6. 200 QPS 边界注记：严禁称为全链路吞吐（必须注明纯路线规划接口 /api/route/plan）
7. 归档与版本检查：严禁在正式目录直接使用未优化旧版本或未归档“初版”文件
"""

import os
import sys
import re
import argparse
from pathlib import Path

# Ensure UTF-8 output on Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import pptx
except ImportError:
    pptx = None

# ANSI colors for terminal output
RED = "\033[91m"
GREEN = "\033[92m"
YELLOW = "\033[93m"
CYAN = "\033[96m"
BOLD = "\033[1m"
RESET = "\033[0m"


class AuditViolation:
    def __init__(self, file_path, line_num, rule_id, severity, matched_text, message, line_content=""):
        self.file_path = str(file_path)
        self.line_num = line_num
        self.rule_id = rule_id
        self.severity = severity  # 'FATAL' or 'WARNING'
        self.matched_text = matched_text
        self.message = message
        self.line_content = line_content.strip()

    def __str__(self):
        color = RED if self.severity == "FATAL" else YELLOW
        location = f"{self.file_path}:{self.line_num}" if self.line_num else self.file_path
        res = [
            f"{color}[{self.severity}] {self.rule_id} at {location}{RESET}",
            f"  {BOLD}Trigger:{RESET} '{self.matched_text}'",
            f"  {BOLD}Reason:{RESET} {self.message}",
        ]
        if self.line_content:
            res.append(f"  {BOLD}Context:{RESET} {self.line_content[:150]}")
        return "\n".join(res)


class MaterialAuditor:
    def __init__(self, root_dir, verbose=False):
        self.root_dir = Path(root_dir).resolve()
        self.verbose = verbose
        self.violations = []
        self.scanned_files = 0

    def is_negative_or_forbidden_spec(self, line: str) -> bool:
        """检查行是否是在声明禁用词、规则说明、或者否定表述"""
        negation_markers = [
            "杜绝", "严禁", "禁", "非", "禁用", "违规", "不采用", "而非", "封存",
            "forbidden", "forbidden_wordings", "当前表述", "历史版本", "旧版", "老草稿",
            "已废弃", "违规表述", "不要写", "禁止", "不一致", "错误表述", "修正前",
            "废除", "替代", "修正为", "早期", "历史", "旧", "淘汰", "草稿", "评委质询"
        ]
        if any(m in line for m in negation_markers):
            return True
        # Markdown 表格中如果是“违规表述”列
        if re.search(r'\|\s*(违规|禁用|错误).*?\|', line):
            return True
        return False

    def extract_text_from_file(self, file_path: Path):
        """返回 [(line_num, line_text)] 列表"""
        ext = file_path.suffix.lower()
        lines = []

        if ext in [".md", ".tex", ".txt", ".yaml", ".yml", ".json"]:
            try:
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    for idx, line in enumerate(f, 1):
                        lines.append((idx, line))
            except Exception as e:
                if self.verbose:
                    print(f"Error reading {file_path}: {e}")

        elif ext == ".pdf":
            if pypdf is None:
                print(f"{YELLOW}Warning: pypdf not installed, skipping PDF content extraction for {file_path}{RESET}")
                return lines
            try:
                reader = pypdf.PdfReader(str(file_path))
                for page_num, page in enumerate(reader.pages, 1):
                    text = page.extract_text() or ""
                    for sub_idx, sub_line in enumerate(text.splitlines(), 1):
                        lines.append((f"Page {page_num}:{sub_idx}", sub_line))
            except Exception as e:
                if self.verbose:
                    print(f"Error reading PDF {file_path}: {e}")

        elif ext == ".pptx":
            if pptx is None:
                print(f"{YELLOW}Warning: python-pptx not installed, skipping PPTX extraction for {file_path}{RESET}")
                return lines
            try:
                prs = pptx.Presentation(str(file_path))
                for slide_num, slide in enumerate(prs.slides, 1):
                    slide_texts = []
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            slide_texts.append(shape.text_frame.text)
                        if shape.has_table:
                            for row in shape.table.rows:
                                for cell in row.cells:
                                    slide_texts.append(cell.text)
                    for text_block in slide_texts:
                        for sub_idx, sub_line in enumerate(text_block.splitlines(), 1):
                            if sub_line.strip():
                                lines.append((f"Slide {slide_num}:{sub_idx}", sub_line))
            except Exception as e:
                if self.verbose:
                    print(f"Error reading PPTX {file_path}: {e}")

        return lines

    def check_file(self, file_path: Path):
        self.scanned_files += 1
        lines = self.extract_text_from_file(file_path)
        is_yaml_spec = file_path.suffix.lower() in [".yaml", ".yml"]
        is_tracker = "remediation_issue_tracker" in file_path.name
        in_forbidden_table = False

        for line_num, line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            # 跟踪 Markdown 表格是否包含“违规/禁用”表头
            if line_str.startswith("|"):
                if any(h in line_str for h in ["违规", "禁用", "错误表述", "错误口径", "forbidden"]):
                    in_forbidden_table = True
            else:
                in_forbidden_table = False

            # 如果是在规格说明文档里的“违规词定义区”或禁用词表格区，跳过常规正文匹配
            if in_forbidden_table or ((is_yaml_spec or is_tracker) and self.is_negative_or_forbidden_spec(line_str)):
                continue
            if is_tracker:
                continue

            # 1. 严禁虚构数据 98.3% / 97.5%
            for bad_rate in ["98.3%", "97.5%"]:
                if bad_rate in line_str:
                    if not self.is_negative_or_forbidden_spec(line_str):
                        self.violations.append(AuditViolation(
                            file_path, line_num, "REDLINE-01-FABRICATED-RATE", "FATAL",
                            bad_rate, "严禁使用未经复现的虚构数值，基线必须为 96.0%，受控消融观察值为 98.0%", line_str
                        ))

            # 2. 路网规模一致性：允许“30 条道路边”，禁止无边界限定的“30 条道路”
            bad_roads = ["23+30", "30+条道路", "30+ 条道路", "30+景点", "30+ 景点"]
            for bad_road in bad_roads:
                if bad_road in line_str:
                    if not self.is_negative_or_forbidden_spec(line_str):
                        self.violations.append(AuditViolation(
                            file_path, line_num, "REDLINE-02-ROAD-COUNT", "FATAL",
                            bad_road, "当前路网资产必须表述为 23 个路线节点、30 条当前登记道路边；31 条仅为待补齐核验的目标口径", line_str
                        ))
            if re.search(r"30\s*条道路(?!边)", line_str) and not self.is_negative_or_forbidden_spec(line_str):
                self.violations.append(AuditViolation(
                    file_path, line_num, "REDLINE-02-ROAD-COUNT", "FATAL",
                    "30 条道路", "当前路网资产必须表述为 23 个路线节点、30 条当前登记道路边；31 条仅为待补齐核验的目标口径", line_str
                ))

            # 3. PPT 页数规范（严禁 13 页）
            bad_slides = ["13页", "13 页"]
            for bad_slide in bad_slides:
                if bad_slide in line_str:
                    if not self.is_negative_or_forbidden_spec(line_str):
                        # 允许表达“Slide 13页”或“第 13 页”讨论第13张幻灯片
                        if not re.search(r'(第\s*13\s*页|Slide\s*13|P13)', line_str):
                            self.violations.append(AuditViolation(
                                file_path, line_num, "REDLINE-03-SLIDE-COUNT", "FATAL",
                                bad_slide, "汇报 PPT 必须统一为 14 页金牌标准版，严禁使用 13 页版本", line_str
                            ))

            # 4. 英雄案例时间线规范（严禁 13:00 / 180分钟 / 3小时 / 13:45 检票口）
            if "13:45" in line_str and "交付时间" not in line_str:
                if not self.is_negative_or_forbidden_spec(line_str):
                    self.violations.append(AuditViolation(
                        file_path, line_num, "REDLINE-04-HERO-TIMELINE", "FATAL",
                        "13:45", "英雄案例抵达梵宫检票口时刻必须严格统一为 13:40（前置 20 分钟排队缓冲）", line_str
                    ))

            # 检查英雄案例是否被缩水为 3小时/180分钟
            if re.search(r'英雄案例.*?(3\s*小时|180\s*分钟|13:00)', line_str) or \
               re.search(r'(3\s*小时|180\s*分钟).*?英雄案例', line_str):
                if not self.is_negative_or_forbidden_spec(line_str):
                    self.violations.append(AuditViolation(
                        file_path, line_num, "REDLINE-04-HERO-DURATION", "FATAL",
                        line_str, "英雄案例总时长必须严格为 300 分钟（5 小时），11:00 出发，16:00 返回正门", line_str
                    ))

            # 5. 98% / 98.0% 观察值限定检查：严禁将 98% 称为基线
            if re.search(r'98(\.0)?%', line_str):
                if not self.is_negative_or_forbidden_spec(line_str):
                    # 如果把 98% 说成基线（且未指明 96% 是基线）
                    is_called_baseline = (
                        re.search(r'(基线|baseline|基准).*?(准确率|通过率)?\s*(为|达到|是|:)?\s*98(\.0)?%', line_str, re.IGNORECASE) or
                        re.search(r'98(\.0)?%\s*(的)?(基线|baseline|基准)', line_str, re.IGNORECASE)
                    )
                    has_96_baseline = "96" in line_str and "基" in line_str
                    has_ablation_qualifier = any(q in line_str for q in ["消融", "观察值", "Full Retrieval", "提升至", "跃迁"])

                    if is_called_baseline and not (has_96_baseline and has_ablation_qualifier):
                        self.violations.append(AuditViolation(
                            file_path, line_num, "REDLINE-05-ABLATION-OBSERVATION", "FATAL",
                            line_str, "98.0% 为受控消融实验中 Full Retrieval 的观察值，权威基线必须明确为 96.0%", line_str
                        ))

            # 6. 200 QPS / 202 QPS 边界注记检查：严禁称为全链路或端到端系统吞吐
            if re.search(r'20[02](\.4)?\s*QPS', line_str, re.IGNORECASE):
                if not self.is_negative_or_forbidden_spec(line_str):
                    is_called_full_chain = (
                        re.search(r'(全链路|端到端|系统综合|系统整体).*?(20[02])', line_str) or
                        re.search(r'20[02].*?(全链路|端到端|系统综合|系统整体)', line_str)
                    )
                    is_negated_full = any(neg in line_str for neg in ["不含", "不包含", "非全链路", "纯路线", "纯规划", "不包括"])
                    if is_called_full_chain and not is_negated_full:
                        self.violations.append(AuditViolation(
                            file_path, line_num, "REDLINE-06-QPS-BOUNDARY", "FATAL",
                            line_str, "200 QPS 实测值必须声明为单台 4 核 8G 边缘节点‘纯路线规划接口 (/api/route/plan)’吞吐，不含 LLM/ASR/TTS", line_str
                        ))

            # Evidence-boundary language: do not claim zero hallucinations or
            # perfect real-world execution when the authoritative baseline is 96%.
            overclaims = [
                "知识零幻觉",
                "根除事实性幻觉",
                "根除幻觉",
                "全程精准履约",
                "决定性贡献",
                "绝对真实可信",
            ]
            if any(term in line_str for term in overclaims):
                if not self.is_negative_or_forbidden_spec(line_str):
                    matched = next(term for term in overclaims if term in line_str)
                    self.violations.append(AuditViolation(
                        file_path, line_num, "REDLINE-08-EVIDENCE-OVERCLAIM", "FATAL",
                        matched, "关键材料必须区分 96% 权威基线、98% 单次观察值、结构化约束测试和真实运营效果", line_str
                    ))

            # Current retrieval architecture is three-way parallel; Router is
            # explanatory metadata and must not be described as disabling branches.
            if "决定检索通道" in line_str or "不同问题 → 不同通道" in line_str:
                if not self.is_negative_or_forbidden_spec(line_str):
                    self.violations.append(AuditViolation(
                        file_path, line_num, "REDLINE-09-ROUTER-SEMANTICS", "FATAL",
                        line_str, "当前口径是 Promise.all 三路并行召回；问题分析只能用于 Trace 标注与结果解释", line_str
                    ))

            # The hero case has one frozen intermediate timeline.
            if "12:40" in line_str:
                if not self.is_negative_or_forbidden_spec(line_str):
                    self.violations.append(AuditViolation(
                        file_path, line_num, "REDLINE-10-HERO-TIMELINE-DRIFT", "FATAL",
                        "12:40", "英雄案例中间节点统一为 12:00~12:20 演出和 12:30 动态纠偏", line_str
                    ))

    def audit_directory(self, target_dir: Path, excludes=None):
        if excludes is None:
            excludes = ["历史版本归档", "node_modules", ".git", ".pytest_cache", ".worktrees"]

        target_dir = Path(target_dir).resolve()
        if not target_dir.exists():
            print(f"{YELLOW}Warning: Directory does not exist: {target_dir}{RESET}")
            return

        for root, dirs, files in os.walk(target_dir):
            # 过滤排除目录
            dirs[:] = [d for d in dirs if not any(ex in d for ex in excludes)]

            for file in files:
                ext = Path(file).suffix.lower()
                if ext in [".md", ".tex", ".pptx", ".pdf"]:
                    # 检查文件名本身是否违规（如残留“初版”）
                    file_path = Path(root) / file
                    if "初版" in file:
                        self.violations.append(AuditViolation(
                            file_path, None, "REDLINE-07-FILENAME-UNOPTIMIZED", "FATAL",
                            file, "正式汇报目录中严禁包含‘初版’未优化文件，必须使用标准命名的正式交付件"
                        ))
                    self.check_file(file_path)
                    if ext == ".pptx" and pptx is not None:
                        try:
                            prs = pptx.Presentation(str(file_path))
                            if len(prs.slides) != 14:
                                self.violations.append(AuditViolation(
                                    file_path, None, "REDLINE-11-PPTX-SLIDE-COUNT", "FATAL",
                                    str(len(prs.slides)), "冻结 PPTX 必须正好包含 14 页", str(len(prs.slides))
                                ))
                        except Exception as e:
                            self.violations.append(AuditViolation(
                                file_path, None, "REDLINE-11-PPTX-READ", "FATAL",
                                str(e), "PPTX 必须可被 python-pptx 正常读取"
                            ))
                    if ext == ".pdf" and pypdf is not None and "PPT" in file_path.name:
                        try:
                            reader = pypdf.PdfReader(str(file_path))
                            if len(reader.pages) != 14:
                                self.violations.append(AuditViolation(
                                    file_path, None, "REDLINE-12-PDF-PAGE-COUNT", "FATAL",
                                    str(len(reader.pages)), "汇报 PPT 导出 PDF 必须正好包含 14 页", str(len(reader.pages))
                                ))
                        except Exception as e:
                            self.violations.append(AuditViolation(
                                file_path, None, "REDLINE-12-PDF-READ", "FATAL",
                                str(e), "汇报 PPT PDF 必须可被 pypdf 正常读取"
                            ))

    def print_report(self):
        print(f"\n{BOLD}======================================================{RESET}")
        print(f"{BOLD}   灵小禅 AIC 全国总决赛材料 95+ 自动化门禁审计报告{RESET}")
        print(f"{BOLD}======================================================{RESET}")
        print(f"扫描文件总数: {self.scanned_files}")
        print(f"违规总项数:   {len(self.violations)}")

        fatal_count = sum(1 for v in self.violations if v.severity == "FATAL")
        warning_count = sum(1 for v in self.violations if v.severity == "WARNING")

        print(f"致命违规 (FATAL):   {RED}{fatal_count}{RESET}")
        print(f"警告提示 (WARNING): {YELLOW}{warning_count}{RESET}\n")

        if self.violations:
            print(f"{RED}{BOLD}--- 违规明细清单 ---{RESET}")
            for v in self.violations:
                print(str(v))
                print("-" * 50)
            print(f"\n{RED}{BOLD}AUDIT FAILED: 存在 {fatal_count} 项致命红线违规，请依据规范立即修复！{RESET}\n")
            return False
        else:
            print(f"{GREEN}{BOLD}AUDIT PASSED: All competition materials conform to 95+ canonical standards.{RESET}\n")
            return True


def main():
    parser = argparse.ArgumentParser(description="灵小禅 95+ 竞赛材料自动化审计工具")
    parser.add_argument("--target-dir", default=None, help="指定审计的目标目录")
    parser.add_argument("--release-dir", default=None, help="指定发布的 release 目录进行独立审计")
    parser.add_argument("--verbose", action="store_true", help="显示详细读取信息")
    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parent.parent
    auditor = MaterialAuditor(repo_root, verbose=args.verbose)

    if args.release_dir:
        release_path = Path(args.release_dir)
        if not release_path.is_absolute():
            release_path = repo_root / release_path
        print(f"{CYAN}>>> 正在审计独立发布包目录: {release_path}{RESET}")
        auditor.audit_directory(release_path, excludes=[])
    elif args.target_dir:
        target_path = Path(args.target_dir)
        if not target_path.is_absolute():
            target_path = repo_root / target_path
        print(f"{CYAN}>>> 正在审计指定目录: {target_path}{RESET}")
        auditor.audit_directory(target_path)
    else:
        # 默认审计两大核心材料包目录
        doc_dir = repo_root / "竞赛汇报与文档材料包"
        comp_docs = repo_root / "docs" / "competition"
        print(f"{CYAN}>>> 正在审计核心材料目录: {doc_dir}{RESET}")
        auditor.audit_directory(doc_dir)
        print(f"{CYAN}>>> 正在审计文档规范目录: {comp_docs}{RESET}")
        auditor.audit_directory(comp_docs)

    passed = auditor.print_report()
    sys.exit(0 if passed else 1)


if __name__ == "__main__":
    main()
